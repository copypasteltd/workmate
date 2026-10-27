"""Generate independent XLSX interoperability fixtures using openpyxl.

Only writes the explicit caller-owned directory. No network or Excel process.
"""
from datetime import datetime
from pathlib import Path
from decimal import Decimal
import json
import sys

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900


def main():
    output = Path(sys.argv[1]).resolve(strict=True)
    for system in ("1900", "1904"):
        book = openpyxl.Workbook()
        book.epoch = CALENDAR_MAC_1904 if system == "1904" else CALENDAR_WINDOWS_1900
        book.active.title = "Cover"
        book.active["A1"] = "Independent workbook fixture; synthetic data"
        sheet = book.create_sheet("香港 Campaign facts")
        sheet.merge_cells("A1:C1")
        sheet["A1"] = "Campaign performance — values only"
        headers = ["date", "campaign", "channel", "impressions", "clicks", "spend", "conversions", "revenue"]
        for column, value in enumerate(headers, 1):
            sheet.cell(3, column, value).font = Font(bold=True)
        rows = [
            [datetime(2026, 10, 1), "香港 Launch\nCampaign", "Google Ads", 1000, 100, Decimal("100.125"), Decimal("4.5"), Decimal("450.875")],
            [datetime(2026, 10, 2), "Values 0", "Google Ads", 2, 1, Decimal("0.000001"), 0, Decimal("0.00125")],
        ]
        for row, values in enumerate(rows, 7):
            for column, value in enumerate(values, 1):
                sheet.cell(row, column, value)
            sheet.cell(row, 1).number_format = "yyyy-mm-dd"
            sheet.cell(row, 6).number_format = '"HK$"#,##0.00000000'
        sheet.row_dimensions[8].hidden = True
        book.create_sheet("Archive").sheet_state = "hidden"
        book.active = sheet
        book.save(output / f"openpyxl-{system}.xlsx")
        if system == "1900":
            sheet["F8"] = Decimal("0.0000001")
            book.save(output / "openpyxl-invalid-precision.xlsx")
            sheet["F8"] = Decimal("0.000001")
            sheet["F7"] = "=SUM(1,2)"
            book.save(output / "openpyxl-formula.xlsx")
        book.close()
    (output / "generator.json").write_text(json.dumps({"generator": "openpyxl", "version": openpyxl.__version__}) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
